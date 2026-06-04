from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from app.models.lesson import ILAWWeeklyPlan


def _set_cell_bg(cell, hex_color: str):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _header_cell(cell, text: str, bold=True, font_size=10, bg="1F3864", fg="FFFFFF"):
    _set_cell_bg(cell, bg)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor.from_string(fg)


def _body_cell(cell, text: str, bold=False, font_size=9):
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)


def export_weekly_plan_docx(plan: ILAWWeeklyPlan) -> bytes:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(1.5)

    # ── Title ─────────────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("FIRST TRIMESTER DAILY LESSON LOG")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("S.Y. 2026-2027").font.size = Pt(10)

    doc.add_paragraph()

    # ── Header info table ─────────────────────────────────────────────────────
    info_table = doc.add_table(rows=2, cols=4)
    info_table.style = "Table Grid"

    labels = ["Name of Teacher", "Learning Area", "Grade Level & Section", "Week"]
    values = [
        plan.teacher_name or "___________________",
        plan.subject,
        plan.grade_level,
        plan.week_label,
    ]

    for i, (label, value) in enumerate(zip(labels, values)):
        info_table.rows[0].cells[i].text = label
        info_table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        info_table.rows[0].cells[i].paragraphs[0].runs[0].font.size = Pt(9)
        info_table.rows[1].cells[i].text = value
        info_table.rows[1].cells[i].paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()

    # ── Learning Competency ───────────────────────────────────────────────────
    lc_table = doc.add_table(rows=1, cols=2)
    lc_table.style = "Table Grid"
    _header_cell(lc_table.rows[0].cells[0], "LEARNING COMPETENCY\n(Kasanayang Pampagkatuto)", font_size=9)
    lc_table.rows[0].cells[0].width = Inches(2)
    _body_cell(lc_table.rows[0].cells[1], plan.learning_competency, font_size=9)

    doc.add_paragraph()

    # ── Learning Objectives: 4 sessions × 3 domains ───────────────────────────
    doc.add_paragraph("LEARNING OBJECTIVES", style="Heading 3")
    obj_table = doc.add_table(rows=5, cols=5)
    obj_table.style = "Table Grid"

    _header_cell(obj_table.rows[0].cells[0], "Domain")
    for i, s in enumerate(plan.sessions):
        _header_cell(obj_table.rows[0].cells[i + 1], f"Session {s.day}\n{s.title}")

    domains = [("Cognitive", "cognitive"), ("Psychomotor", "psychomotor"), ("Affective", "affective")]
    for row_i, (domain_label, domain_key) in enumerate(domains):
        _header_cell(obj_table.rows[row_i + 1].cells[0], domain_label, bg="2E75B6")
        for col_i, session in enumerate(plan.sessions):
            text = getattr(session.objectives, domain_key, "")
            _body_cell(obj_table.rows[row_i + 1].cells[col_i + 1], text)

    # Learning Resources row
    _header_cell(obj_table.rows[4].cells[0], "Learning Resources", bg="2E75B6")
    merged = obj_table.rows[4].cells[1]
    merged.merge(obj_table.rows[4].cells[4])
    _body_cell(merged, plan.learning_resources)

    doc.add_paragraph()

    # ── 4 Session plans ───────────────────────────────────────────────────────
    for session in plan.sessions:
        doc.add_paragraph(
            f"SESSION {session.day}: {session.title.upper()}  ({session.session_type})",
            style="Heading 3"
        )

        s_table = doc.add_table(rows=5, cols=2)
        s_table.style = "Table Grid"
        s_table.columns[0].width = Inches(1.8)

        rows_data = [
            ("PRE-LESSON\n(Warm-up / Pampagana)", session.pre_lesson),
            ("LEARNING EXPERIENCE\n(Karanasan sa Pagkatuto)", session.learning_experience),
            ("FORMATIVE ASSESSMENT\n(Pagtataya)", session.formative_assessment),
            ("INTEGRATION OPPORTUNITIES", session.integration_opportunities or "—"),
        ]

        for r_i, (label, content) in enumerate(rows_data):
            _header_cell(s_table.rows[r_i].cells[0], label, font_size=8, bg="2E75B6")
            _body_cell(s_table.rows[r_i].cells[1], content)

        # Notes row
        _header_cell(s_table.rows[4].cells[0], "NOTES FOR MYSELF\n(Reflection)", font_size=8, bg="70AD47")
        _body_cell(s_table.rows[4].cells[1], "_______________________________________________\n_______________________________________________")

        doc.add_paragraph()

    # ── Ways Forward ──────────────────────────────────────────────────────────
    doc.add_paragraph("WAYS FORWARD (W)", style="Heading 3")
    wf_table = doc.add_table(rows=4, cols=2)
    wf_table.style = "Table Grid"
    wf_table.columns[0].width = Inches(1.8)

    wf_rows = [
        ("REMEDIATION\n(Para sa mga nahihirapan)", plan.ways_forward.get("remediation", "")),
        ("ENRICHMENT\n(Para sa mga advanced)", plan.ways_forward.get("enrichment", "")),
        ("EXTENDED LEARNING\n(Take-home / Pamilya)", plan.ways_forward.get("extended_learning", "")),
        ("TEACHER REFLECTION\n(Pagninilay ng Guro)", plan.ways_forward.get("teacher_reflection", "")),
    ]
    for r_i, (label, content) in enumerate(wf_rows):
        _header_cell(wf_table.rows[r_i].cells[0], label, font_size=8, bg="833C00", fg="FFFFFF")
        _body_cell(wf_table.rows[r_i].cells[1], content)

    doc.add_paragraph()

    # ── Integrated Activities ─────────────────────────────────────────────────
    if plan.integrated_activities:
        doc.add_paragraph("SUBJECT INTEGRATION", style="Heading 3")
        for subj, activity in plan.integrated_activities.items():
            p = doc.add_paragraph()
            p.add_run(f"{subj}: ").bold = True
            p.add_run(activity).font.size = Pt(9)

        doc.add_paragraph()

    # ── AI Declaration ────────────────────────────────────────────────────────
    ai_table = doc.add_table(rows=1, cols=2)
    ai_table.style = "Table Grid"
    _header_cell(ai_table.rows[0].cells[0], "DECLARATION OF AI USE\n(DO 3, s. 2026)", font_size=8, bg="595959")
    _body_cell(ai_table.rows[0].cells[1], plan.ai_declaration, font_size=8)

    doc.add_paragraph()

    # ── Signature block ───────────────────────────────────────────────────────
    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.style = "Table Grid"
    for label, col in [("Prepared by:", 0), ("Checked by (Head Teacher):", 1), ("Noted by (School Head):", 2)]:
        sig_table.rows[0].cells[col].text = label
        sig_table.rows[0].cells[col].paragraphs[0].runs[0].bold = True
        sig_table.rows[0].cells[col].paragraphs[0].runs[0].font.size = Pt(8)
        sig_table.rows[1].cells[col].text = "\n\n_______________________"
        sig_table.rows[1].cells[col].paragraphs[0].runs[0].font.size = Pt(9)

    # Save to buffer
    import io
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
